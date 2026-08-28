from datetime import date,datetime,timezone
import logging
from decimal import Decimal
from io import BytesIO
from typing import Literal
from uuid import UUID
from zoneinfo import ZoneInfo
from fastapi import APIRouter,Depends,File,HTTPException,Query,UploadFile
from fastapi.responses import FileResponse,StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel,Field
from sqlalchemy import func,select,text
from sqlalchemy.orm import Session,selectinload
from app.config import get_settings
from app.database import get_db
from app.models import *
from app.repositories.expenses import ExpenseFilters,ExpenseRepository
from app.services.finance import distribute_evenly,expense_totals,invoice_totals
from app.services.expense_import import import_expenses_excel
from app.services.ocr import get_provider
from app.services.ocr.base import OCRResult
from app.services.storage import save_bytes
from app.services.notifications import accounting_message,send_accounting_email
router=APIRouter()
logger=logging.getLogger(__name__)
class PartnerIn(BaseModel): name:str=Field(min_length=1,max_length=255); comment:str|None=None; counterparty_ids:list[UUID]|None=None
class CounterpartyIn(BaseModel): partner_id:UUID|None=None; full_name:str; short_name:str|None=None; entity_type:str; inn:str|None=None; kpp:str|None=None; comment:str|None=None
class AllocationIn(BaseModel): store_id:UUID; amount:Decimal=Field(default=Decimal(0),ge=0)
class ExpenseIn(BaseModel): partner_id:UUID; counterparty_id:UUID; service_name:str; expense_month:int=Field(ge=1,le=12); expense_year:int=Field(ge=2000,le=2200); contract_number:str|None=None; contract_date:date|None=None; comment:str|None=None; allocations:list[AllocationIn]=Field(default_factory=list); tag_ids:list[UUID]=Field(default_factory=list)
class ExpenseBulkUpdateIn(BaseModel): expense_ids:list[UUID]=Field(min_length=1); partner_id:UUID|None=None; counterparty_id:UUID|None=None; tag_ids:list[UUID]|None=None
class InvoiceIn(BaseModel): invoice_number:str; invoice_date:date; amount:Decimal=Field(ge=0); vat_amount:Decimal|None=Field(default=None,ge=0); comment:str|None=None; allow_duplicate:bool=False
class PaymentIn(BaseModel): payment_date:date; amount:Decimal=Field(ge=0); comment:str|None=None
class StoreIn(BaseModel): name:str=Field(min_length=1,max_length=255); address:str|None=None; comment:str|None=None
class StorePresetIn(BaseModel): name:str=Field(min_length=1,max_length=255); store_ids:list[UUID]=Field(min_length=1)
class TagIn(BaseModel): name:str=Field(min_length=1,max_length=100)
class AISettingsIn(BaseModel): enabled:bool; model:str=Field(min_length=1,max_length=100); api_key:str|None=None
class AILogStatusIn(BaseModel): status:str=Field(pattern="^(new|prompt_copied|in_progress|fixed|no_fix_required)$")
@router.get("/health")
def health(db:Session=Depends(get_db)): db.execute(text("select 1")); return {"status":"ok","database":"ok"}
@router.get("/dashboard")
def dashboard(period:str=Query("month",pattern="^(month|quarter|year)$"),tag_ids:list[UUID]=Query(default=[]),store_ids:list[UUID]=Query(default=[]),db:Session=Depends(get_db)):
 today=datetime.now(ZoneInfo(get_settings().app_timezone)).date(); start_month=today.month if period=="month" else ((today.month-1)//3)*3+1 if period=="quarter" else 1
 q=select(Expense).where(Expense.deleted_at.is_(None),Expense.expense_year==today.year,Expense.expense_month>=start_month,Expense.expense_month<=(start_month+2 if period=="quarter" else today.month if period=="month" else 12)).options(selectinload(Expense.invoices).selectinload(Invoice.payments),selectinload(Expense.tags))
 if tag_ids: q=q.where(Expense.tags.any(Tag.id.in_(tag_ids)))
 if store_ids: q=q.where(Expense.allocations.any(Allocation.store_id.in_(store_ids)))
 items=db.scalars(q).unique().all(); invoice_total=paid_total=Decimal(0); tag_totals={}
 for expense in items:
  expense_total=Decimal(0)
  for invoice in expense.invoices:
   if invoice.deleted_at: continue
   invoice_total+=invoice.amount; expense_total+=invoice.amount; paid_total+=sum((payment.amount for payment in invoice.payments if not payment.deleted_at),Decimal(0))
  labels=[tag.name for tag in expense.tags] or ["Без тега"]
  for label in labels:
   current=tag_totals.setdefault(label,{"amount":Decimal(0),"expense_count":0}); current["amount"]+=expense_total; current["expense_count"]+=1
 return {"invoice_total":invoice_total,"paid_total":paid_total,"remaining_total":invoice_total-paid_total,"expense_count":len(items),"period":period,"tag_totals":[{"tag":tag,"amount":value["amount"],"expense_count":value["expense_count"]} for tag,value in sorted(tag_totals.items(),key=lambda item:item[1]["amount"],reverse=True)]}
@router.get("/partners")
def partners(search:str|None=None,db:Session=Depends(get_db)):
 q=select(Partner).where(Partner.deleted_at.is_(None)); q=q.where(Partner.name.ilike(f"%{search}%")) if search else q
 return db.scalars(q.order_by(Partner.name)).all()
@router.post("/partners",status_code=201)
def create_partner(data:PartnerIn,db:Session=Depends(get_db)):
 x=Partner(**data.model_dump(exclude={"counterparty_ids"})); db.add(x); db.flush()
 for item_id in data.counterparty_ids or []:
  item=db.get(Counterparty,item_id)
  if item and not item.deleted_at: item.partner_id=x.id
 db.commit(); return x
@router.get("/partners/{item_id}")
def partner_detail(item_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Partner,item_id)
 if not x or x.deleted_at: raise HTTPException(404,"Партнер не найден")
 return {"id":x.id,"name":x.name,"comment":x.comment,"counterparties":[{"id":item.id,"full_name":item.full_name,"inn":item.inn,"kpp":item.kpp} for item in x.counterparties if not item.deleted_at]}
@router.put("/partners/{item_id}")
def update_partner(item_id:UUID,data:PartnerIn,db:Session=Depends(get_db)):
 x=db.get(Partner,item_id)
 if not x or x.deleted_at: raise HTTPException(404,"Партнер не найден")
 for key,value in data.model_dump(exclude={"counterparty_ids"}).items(): setattr(x,key,value)
 if data.counterparty_ids is not None:
  selected=set(data.counterparty_ids)
  for item in db.scalars(select(Counterparty).where(Counterparty.deleted_at.is_(None))).all():
   if item.partner_id==x.id and item.id not in selected: item.partner_id=None
   elif item.id in selected: item.partner_id=x.id
 db.commit(); db.refresh(x); return x
@router.delete("/partners/{item_id}",status_code=204)
def archive_partner(item_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Partner,item_id)
 if not x: raise HTTPException(404,"Партнер не найден")
 x.deleted_at=datetime.now(timezone.utc); db.commit()
@router.get("/counterparties")
def counterparties(search:str|None=None,inn:str|None=None,db:Session=Depends(get_db)):
 q=select(Counterparty).where(Counterparty.deleted_at.is_(None)); q=q.where(Counterparty.inn==inn) if inn else q; q=q.where(Counterparty.full_name.ilike(f"%{search}%")) if search else q; return db.scalars(q).all()
@router.post("/counterparties",status_code=201)
def create_counterparty(data:CounterpartyIn,db:Session=Depends(get_db)):
 if data.partner_id and not db.get(Partner,data.partner_id): raise HTTPException(422,"Партнер не найден")
 x=Counterparty(**data.model_dump()); db.add(x); db.commit(); return x
@router.get("/counterparties/{item_id}")
def counterparty_detail(item_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Counterparty,item_id)
 if not x or x.deleted_at: raise HTTPException(404,"Контрагент не найден")
 return {"id":x.id,"partner_id":x.partner_id,"full_name":x.full_name,"short_name":x.short_name,"entity_type":x.entity_type,"inn":x.inn,"kpp":x.kpp,"comment":x.comment}
@router.put("/counterparties/{item_id}")
def update_counterparty(item_id:UUID,data:CounterpartyIn,db:Session=Depends(get_db)):
 x=db.get(Counterparty,item_id)
 if not x or x.deleted_at: raise HTTPException(404,"Контрагент не найден")
 for key,value in data.model_dump().items(): setattr(x,key,value)
 db.commit(); db.refresh(x); return x
@router.delete("/counterparties/{item_id}",status_code=204)
def archive_counterparty(item_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Counterparty,item_id)
 if not x: raise HTTPException(404,"Контрагент не найден")
 x.deleted_at=datetime.now(timezone.utc); db.commit()
@router.get("/stores")
def stores(db:Session=Depends(get_db)):
 return db.scalars(select(Store).where(Store.is_active.is_(True),Store.name!="Общий маркетинг").order_by(Store.name)).all()
@router.post("/stores",status_code=201)
def create_store(data:StoreIn,db:Session=Depends(get_db)):
 x=Store(**data.model_dump()); db.add(x); db.commit(); return x
@router.put("/stores/{item_id}")
def update_store(item_id:UUID,data:StoreIn,db:Session=Depends(get_db)):
 x=db.get(Store,item_id)
 if not x: raise HTTPException(404,"Магазин не найден")
 for key,value in data.model_dump().items(): setattr(x,key,value)
 db.commit(); db.refresh(x); return x
@router.delete("/stores/{item_id}",status_code=204)
def archive_store(item_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Store,item_id)
 if not x or x.is_system: raise HTTPException(422,"Системный магазин удалить нельзя")
 x.is_active=False; db.commit()
@router.get("/tags")
def tags(db:Session=Depends(get_db)):
 return db.scalars(select(Tag).order_by(Tag.name)).all()
@router.post("/tags",status_code=201)
def create_tag(data:TagIn,db:Session=Depends(get_db)):
 x=Tag(**data.model_dump()); db.add(x); db.commit(); return x
@router.put("/tags/{item_id}")
def update_tag(item_id:UUID,data:TagIn,db:Session=Depends(get_db)):
 x=db.get(Tag,item_id)
 if not x: raise HTTPException(404,"Тег не найден")
 x.name=data.name; db.commit(); db.refresh(x); return x
@router.delete("/tags/{item_id}",status_code=204)
def delete_tag(item_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Tag,item_id)
 if not x: raise HTTPException(404,"Тег не найден")
 for expense in db.scalars(select(Expense).where(Expense.tags.any(Tag.id==item_id))).all(): expense.tags.remove(x)
 db.delete(x); db.commit()
@router.get("/expenses")
def expenses(page:int=Query(1,ge=1),page_size:int=Query(25,ge=25,le=100),search:str|None=None,db:Session=Depends(get_db)):
 items,total=ExpenseRepository(db).list(page,page_size,search); out=[]; documents_by_expense={x.id:set() for x in items}
 if items:
  for expense_id,document_type in db.execute(select(Document.expense_id,Document.document_type).where(Document.expense_id.in_(documents_by_expense),Document.deleted_at.is_(None))): documents_by_expense[expense_id].add(document_type)
 for x in items:
  it,paid,remaining=expense_totals([(i.amount,[p.amount for p in i.payments if not p.deleted_at]) for i in x.invoices if not i.deleted_at]); document_types=documents_by_expense[x.id]; out.append({"id":x.id,"partner":x.partner.name,"counterparty":x.counterparty.full_name,"stores":[a.store.name for a in x.allocations],"tags":[t.name for t in x.tags],"has_invoice_document":"invoice" in document_types,"has_closing_document":"closing" in document_types,"service_name":x.service_name,"period":f"{x.expense_month:02d}.{x.expense_year}","invoice_total":it,"paid_total":paid,"remaining_total":remaining,"updated_at":x.updated_at})
 return {"items":out,"total":total,"page":page,"page_size":page_size}
@router.get("/expenses/ids")
def expense_ids(search:str|None=None,period:str|None=Query(None,pattern=r"^(0[1-9]|1[0-2])\.(20\d{2}|21\d{2})$"),payment_status:Literal["all","paid","unpaid"]="all",partner_ids:list[UUID]=Query(default=[]),counterparty_ids:list[UUID]=Query(default=[]),store_ids:list[UUID]=Query(default=[]),tag_ids:list[UUID]=Query(default=[]),amount_from:Decimal|None=Query(None,ge=0),amount_to:Decimal|None=Query(None,ge=0),invoice_date_from:date|None=None,invoice_date_to:date|None=None,invoice_document:Literal["all","yes","no","cash"]="all",closing_document:Literal["all","yes","no"]="all",db:Session=Depends(get_db)):
 if amount_from is not None and amount_to is not None and amount_from>amount_to: raise HTTPException(422,"Минимальная сумма не может быть больше максимальной")
 if invoice_date_from and invoice_date_to and invoice_date_from>invoice_date_to: raise HTTPException(422,"Дата счета от не может быть позже даты счета до")
 month,year=(map(int,period.split(".")) if period else (None,None))
 filters=ExpenseFilters(search=search,expense_month=month,expense_year=year,payment_status=payment_status,partner_ids=tuple(partner_ids),counterparty_ids=tuple(counterparty_ids),store_ids=tuple(store_ids),tag_ids=tuple(tag_ids),amount_from=amount_from,amount_to=amount_to,invoice_date_from=invoice_date_from,invoice_date_to=invoice_date_to,invoice_document=invoice_document,closing_document=closing_document)
 return {"ids":ExpenseRepository(db).ids(filters)}
@router.post("/expenses",status_code=201)
def create_expense(data:ExpenseIn,db:Session=Depends(get_db)):
 if not db.get(Partner,data.partner_id) or not db.get(Counterparty,data.counterparty_id): raise HTTPException(422,"Партнер или контрагент не найден")
 if any(not db.get(Store,item.store_id) for item in data.allocations): raise HTTPException(422,"Магазин не найден")
 values=data.model_dump(exclude={"allocations","tag_ids"}); x=Expense(**values); x.tags=[db.get(Tag,item_id) for item_id in data.tag_ids if db.get(Tag,item_id)]; db.add(x); db.flush()
 for item in data.allocations: db.add(Allocation(expense_id=x.id,**item.model_dump()))
 if data.ai_log_id:
  ai_log=db.get(AIFallbackLog,data.ai_log_id)
  if ai_log and ai_log.expense_id is None:
   ai_log.expense_id=x.id; ai_log.final_fields={**ai_log.final_fields,"partner":{"id":str(data.partner_id),"value":db.get(Partner,data.partner_id).name,"source":"manual"},"counterparty":{"id":str(data.counterparty_id),"value":db.get(Counterparty,data.counterparty_id).full_name,"source":"manual"},"service_name":{"value":data.service_name,"source":"manual"}}; ai_log.updated_at=datetime.now(timezone.utc)
 db.add(AuditLog(entity_type="expense",entity_id=x.id,action="created",metadata_={},created_at=datetime.now(timezone.utc))); db.commit(); return {"id":x.id}
@router.put("/expenses/bulk/update")
def bulk_update_expenses(data:ExpenseBulkUpdateIn,db:Session=Depends(get_db)):
 fields=data.model_fields_set-{"expense_ids"}
 if not fields: raise HTTPException(422,"Выберите хотя бы одно поле для изменения")
 expenses=db.scalars(select(Expense).where(Expense.id.in_(data.expense_ids),Expense.deleted_at.is_(None))).unique().all()
 if len(expenses)!=len(set(data.expense_ids)): raise HTTPException(404,"Один или несколько расходов не найдены")
 partner=db.get(Partner,data.partner_id) if "partner_id" in fields else None
 counterparty=db.get(Counterparty,data.counterparty_id) if "counterparty_id" in fields else None
 if "partner_id" in fields and (not partner or partner.deleted_at): raise HTTPException(422,"Партнер не найден")
 if "counterparty_id" in fields and (not counterparty or counterparty.deleted_at): raise HTTPException(422,"Контрагент не найден")
 if partner and counterparty and counterparty.partner_id and counterparty.partner_id!=partner.id: raise HTTPException(422,"Контрагент не связан с выбранным партнером")
 selected_tags=None
 if "tag_ids" in fields:
  selected_tags=db.scalars(select(Tag).where(Tag.id.in_(data.tag_ids or []))).all()
  if len(selected_tags)!=len(set(data.tag_ids or [])): raise HTTPException(422,"Один или несколько тегов не найдены")
 for expense in expenses:
  if partner: expense.partner_id=partner.id
  if counterparty:
   expense.counterparty_id=counterparty.id
   if "partner_id" not in fields and counterparty.partner_id: expense.partner_id=counterparty.partner_id
  if selected_tags is not None: expense.tags=list(selected_tags)
  db.add(AuditLog(entity_type="expense",entity_id=expense.id,action="bulk_updated",metadata_={"fields":sorted(fields)},created_at=datetime.now(timezone.utc)))
 db.commit(); return {"updated":len(expenses)}
@router.get("/expenses/{expense_id}")
def expense_detail(expense_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Expense,expense_id)
 if not x or x.deleted_at: raise HTTPException(404,"Расход не найден")
 return {"id":x.id,"partner_id":x.partner_id,"counterparty_id":x.counterparty_id,"service_name":x.service_name,"expense_month":x.expense_month,"expense_year":x.expense_year,"contract_number":x.contract_number,"contract_date":x.contract_date,"comment":x.comment,"allocations":[{"store_id":a.store_id,"store":a.store.name,"amount":a.amount} for a in x.allocations],"tags":[{"id":t.id,"name":t.name} for t in x.tags],"invoices":[{"id":i.id,"invoice_number":i.invoice_number,"invoice_date":i.invoice_date,"amount":i.amount,"payments":[{"id":p.id,"payment_date":p.payment_date,"amount":p.amount,"comment":p.comment} for p in i.payments if not p.deleted_at]} for i in x.invoices if not i.deleted_at],"documents":[{"id":d.id,"document_type":d.document_type,"original_filename":d.original_filename,"mime_type":d.mime_type,"created_at":d.created_at} for d in db.scalars(select(Document).where(Document.expense_id==x.id,Document.deleted_at.is_(None))).all()]}
@router.put("/expenses/{expense_id}")
def update_expense(expense_id:UUID,data:ExpenseIn,db:Session=Depends(get_db)):
 x=db.get(Expense,expense_id)
 if not x or x.deleted_at: raise HTTPException(404,"Расход не найден")
 for key,value in data.model_dump(exclude={"allocations","tag_ids"}).items(): setattr(x,key,value)
 existing={item.store_id:item for item in x.allocations}; selected={item.store_id:item for item in data.allocations}
 for store_id,allocation in list(existing.items()):
  if store_id not in selected: db.delete(allocation)
 for store_id,item in selected.items():
  if store_id in existing: existing[store_id].amount=item.amount
  else: x.allocations.append(Allocation(**item.model_dump()))
 x.tags=[db.get(Tag,item_id) for item_id in data.tag_ids if db.get(Tag,item_id)]
 db.commit(); return {"id":x.id}
@router.delete("/expenses/{expense_id}",status_code=204)
def delete_expense(expense_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Expense,expense_id)
 if not x: raise HTTPException(404,"Расход не найден")
 x.deleted_at=datetime.now(timezone.utc); db.commit()
@router.post("/expenses/{expense_id}/documents",status_code=201)
async def upload_expense_document(expense_id:UUID,document_type:str=Query(pattern="^(invoice|closing)$"),file:UploadFile=File(...),db:Session=Depends(get_db)):
 if not db.get(Expense,expense_id): raise HTTPException(404,"Расход не найден")
 data=await file.read(); s=get_settings()
 try: stored,path,sha=save_bytes(data,file.filename or "document",file.content_type or "",s.upload_dir,s.max_upload_size_mb)
 except ValueError as e: raise HTTPException(422,str(e))
 x=Document(expense_id=expense_id,document_type=document_type,original_filename=file.filename or "document",stored_filename=stored,storage_path=path,mime_type=file.content_type or "",file_size=len(data),sha256=sha,created_at=datetime.now(timezone.utc)); db.add(x); db.commit(); return {"id":x.id}
@router.get("/documents/{document_id}/content")
def document_content(document_id:UUID,db:Session=Depends(get_db)):
 x=db.get(Document,document_id)
 if not x or x.deleted_at: raise HTTPException(404,"Документ не найден")
 return FileResponse(x.storage_path,media_type=x.mime_type,filename=x.original_filename,content_disposition_type="inline")
@router.put("/documents/{document_id}/expense/{expense_id}")
def attach_document(document_id:UUID,expense_id:UUID,db:Session=Depends(get_db)):
 document=db.get(Document,document_id); expense=db.get(Expense,expense_id)
 if not document or document.deleted_at: raise HTTPException(404,"Документ не найден")
 if not expense or expense.deleted_at: raise HTTPException(404,"Расход не найден")
 document.expense_id=expense_id; db.commit(); return {"id":document.id}
@router.post("/expenses/{expense_id}/invoices",status_code=201)
def add_invoice(expense_id:UUID,data:InvoiceIn,db:Session=Depends(get_db)):
 exp=db.get(Expense,expense_id)
 if not exp or exp.deleted_at: raise HTTPException(404,"Расход не найден")
 dup=db.scalar(select(Invoice).join(Expense).join(Counterparty).where(Counterparty.inn==exp.counterparty.inn,func.lower(Invoice.invoice_number)==data.invoice_number.lower(),Invoice.invoice_date==data.invoice_date,Invoice.amount==data.amount,Invoice.deleted_at.is_(None)))
 if dup and not data.allow_duplicate: raise HTTPException(409,detail={"message":"Возможно, этот счет уже существует","invoice_id":str(dup.id),"expense_id":str(dup.expense_id)})
 x=Invoice(expense_id=expense_id,**data.model_dump(exclude={"allow_duplicate"})); db.add(x); db.flush(); db.add(AuditLog(entity_type="invoice",entity_id=x.id,action="created",metadata_={},created_at=datetime.now(timezone.utc)))
 db.commit()
 document=db.scalar(select(Document).where(Document.expense_id==expense_id,Document.document_type=="invoice",Document.deleted_at.is_(None)).order_by(Document.created_at.desc()))
 notification_sent=False
 if document:
  store_names=[allocation.store.name for allocation in exp.allocations]
  message=accounting_message(exp.service_name,x.amount,store_names)
  notification_sent=send_accounting_email(get_settings(),subject=f"Счет на оплату: {exp.service_name}",message=message,attachment_path=document.storage_path,attachment_name=document.original_filename,attachment_mime=document.mime_type)
 return {"id":x.id,"notification_sent":notification_sent}
@router.put("/invoices/{invoice_id}")
def update_invoice(invoice_id:UUID,data:InvoiceIn,db:Session=Depends(get_db)):
 x=db.get(Invoice,invoice_id)
 if not x or x.deleted_at: raise HTTPException(404,"Счет не найден")
 for key,value in data.model_dump(exclude={"allow_duplicate"},exclude_unset=True).items(): setattr(x,key,value)
 db.commit(); return {"id":x.id}
@router.post("/invoices/{invoice_id}/payments",status_code=201)
def add_payment(invoice_id:UUID,data:PaymentIn,db:Session=Depends(get_db)):
 inv=db.get(Invoice,invoice_id)
 if not inv or inv.deleted_at: raise HTTPException(404,"Счет не найден")
 x=Payment(invoice_id=invoice_id,**data.model_dump()); db.add(x); db.flush(); db.add(AuditLog(entity_type="payment",entity_id=x.id,action="created",metadata_={},created_at=datetime.now(timezone.utc))); db.commit(); paid,remaining=invoice_totals(inv.amount,[p.amount for p in inv.payments if not p.deleted_at]); return {"id":x.id,"paid_amount":paid,"remaining_amount":remaining}
@router.put("/payments/{payment_id}")
def update_payment(payment_id:UUID,data:PaymentIn,db:Session=Depends(get_db)):
 x=db.get(Payment,payment_id)
 if not x or x.deleted_at: raise HTTPException(404,"Платеж не найден")
 for key,value in data.model_dump().items(): setattr(x,key,value)
 db.commit(); return {"id":x.id}
def _serialize_ocr(result:OCRResult):
 values={"invoice_number":result.invoice_number,"invoice_date":result.invoice_date,"amount":str(result.invoice_amount) if result.invoice_amount is not None else None,"recipient":result.counterparty_name,"inn":result.inn,"kpp":result.kpp}
 confidence={"invoice_number":result.confidence.get("invoice_number",0),"invoice_date":result.confidence.get("invoice_date",0),"amount":result.confidence.get("invoice_amount",0),"recipient":result.confidence.get("counterparty_name",0),"inn":result.confidence.get("inn",0),"kpp":result.confidence.get("kpp",result.confidence.get("inn",0))}
 return values,confidence
def _match_counterparty(result:OCRResult,db:Session):
 normalized_inn="".join(x for x in (result.inn or "") if x.isdigit())
 active=db.scalars(select(Counterparty).where(Counterparty.deleted_at.is_(None))).all()
 matches=[item for item in active if "".join(x for x in (item.inn or "") if x.isdigit())==normalized_inn] if normalized_inn else []
 if len(matches)==1: return matches[0],True
 if result.counterparty_name and result.confidence.get("counterparty_name",0)>=.85:
  normalized=lambda value:"".join(x for x in value.lower() if x.isalnum())
  target=normalized(result.counterparty_name); names=[x for x in active if normalized(x.full_name)==target]
  if len(names)==1: return names[0],True
 return None,False
def _run_recognition(document:Document,db:Session):
 s=get_settings()
 try: result=get_provider(s.ocr_provider).recognize(document.storage_path,document.mime_type)
 except Exception as error:
  # Техническая причина остаётся в серверном журнале, а пользователь получает
  # понятное сообщение и может продолжить ручное заполнение.
  logger.exception("OCR failed for document %s using provider %s",document.id,s.ocr_provider)
  raise HTTPException(503,"Сервис распознавания временно недоступен. Попробуйте еще раз или заполните данные вручную.") from error
 values,confidence=_serialize_ocr(result); counterparty,matched=_match_counterparty(result,db)
 if matched: confidence["inn"]=max(confidence["inn"],.99); confidence["recipient"]=max(confidence["recipient"],.95)
 recognition=OCRRecognition(document_id=document.id,provider=s.ocr_provider,raw_text=result.raw_text,fields=values,confidence=confidence,blocks=result.blocks,created_at=datetime.now(timezone.utc)); db.add(recognition); db.commit()
 partner=db.get(Partner,counterparty.partner_id) if counterparty and counterparty.partner_id else None
 primary={"partner":{"id":str(partner.id) if partner else None,"value":partner.name if partner else None,"matched":bool(partner),"confidence":.99 if partner else 0,"source":"original"},"counterparty":{"id":str(counterparty.id) if counterparty else None,"value":counterparty.full_name if counterparty else result.counterparty_name,"inn":result.inn,"matched":matched,"confidence":confidence["recipient"],"source":"original"},"service_name":{"value":result.service_name,"confidence":result.confidence.get("service_name",0),"source":"original"},"invoice_number":{"value":result.invoice_number,"confidence":confidence["invoice_number"],"source":"original"},"invoice_date":{"value":result.invoice_date,"confidence":confidence["invoice_date"],"source":"original"},"amount":{"value":str(result.invoice_amount) if result.invoice_amount is not None else None,"confidence":confidence["amount"],"source":"original"}}
 response_fields={"invoice_number":primary["invoice_number"],"invoice_date":primary["invoice_date"],"amount":primary["amount"],"recipient":{"value":primary["counterparty"].get("value"),"confidence":primary["counterparty"].get("confidence",0),"source":"original"},"inn":{"value":result.inn,"confidence":confidence["inn"],"source":"original"},"kpp":{"value":result.kpp,"confidence":confidence["kpp"],"source":"original"},"service_name":primary["service_name"]}
 return {"status":"success","document_id":document.id,"fields":response_fields,"partner":{"matched":bool(partner),"id":partner.id if partner else None,"name":partner.name if partner else None,"suggestion":None,"source":"original"},"counterparty":{"matched":matched,"id":counterparty.id if counterparty else None,"name":counterparty.full_name if counterparty else result.counterparty_name,"suggestion":None if matched else result.counterparty_name,"source":"original"},"raw_text":result.raw_text}
@router.post("/ocr/invoice")
@router.post("/ocr")
async def ocr(file:UploadFile=File(...),db:Session=Depends(get_db)):
 data=await file.read(); s=get_settings()
 try: stored,path,sha=save_bytes(data,file.filename or "document",file.content_type or "",s.upload_dir,s.max_upload_size_mb)
 except ValueError as e: raise HTTPException(422,str(e))
 document=Document(document_type="invoice",original_filename=file.filename or "document",stored_filename=stored,storage_path=path,mime_type=file.content_type or "",file_size=len(data),sha256=sha,created_at=datetime.now(timezone.utc)); db.add(document); db.commit(); db.refresh(document); return _run_recognition(document,db)
@router.post("/documents/{document_id}/recognize")
def recognize_document(document_id:UUID,db:Session=Depends(get_db)):
 document=db.get(Document,document_id)
 if not document or document.deleted_at: raise HTTPException(404,"Документ не найден")
 return _run_recognition(document,db)

@router.get("/export")
def export(db:Session=Depends(get_db)):
 items,_=ExpenseRepository(db).list(1,100,None); wb=Workbook(); ws=wb.active; ws.title="Расходы"; ws.append(["Период","Партнер","Контрагент","ИНН","Услуга","Договор","Счет","Дата счета","Сумма","Оплачено","Остаток"])
 for e in items:
  for i in e.invoices or [None]:
   paid,remain=invoice_totals(i.amount,[p.amount for p in i.payments if not p.deleted_at]) if i else (Decimal(0),Decimal(0)); ws.append([f"{e.expense_month:02d}.{e.expense_year}",e.partner.name,e.counterparty.full_name,e.counterparty.inn,e.service_name,e.contract_number,i.invoice_number if i else "",i.invoice_date if i else "",i.amount if i else 0,paid,remain])
 stream=BytesIO(); wb.save(stream); stream.seek(0); return StreamingResponse(stream,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=expenses.xlsx"})

@router.post("/expenses/import-xlsx")
async def import_expenses(file:UploadFile=File(...),db:Session=Depends(get_db)):
 filename=file.filename or ""
 if not filename.lower().endswith((".xlsx",".xls")):
  raise HTTPException(422,"Выберите файл в формате XLSX или XLS")
 data=await file.read()
 if not data: raise HTTPException(422,"Файл пуст")
 try: return import_expenses_excel(data,filename,db)
 except ValueError as error:
  db.rollback(); raise HTTPException(422,str(error)) from error
